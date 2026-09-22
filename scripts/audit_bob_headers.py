from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/upload/BOB2026ledgerMASTER.xlsx')
wb = load_workbook(path, read_only=True, data_only=False)
for ws in wb.worksheets:
    print(f'=== {ws.title} ===')
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        rows.append(list(row))
    for ri, row in enumerate(rows, 1):
        vals = [(i + 1, v) for i, v in enumerate(row) if v not in (None, '')]
        print(f'row {ri}: {vals}')
    print()
