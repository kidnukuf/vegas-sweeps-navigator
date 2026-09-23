from __future__ import annotations

import json
from pathlib import Path
from openpyxl import load_workbook

root = Path('/home/ubuntu/vegas-sweeps-navigator/work')
source_path = root / 'source_delegate_9_22.xlsx'
master_path = root / 'realnovemberevent_before.json'

wb = load_workbook(source_path, data_only=True, read_only=False)
print('SOURCE_WORKBOOK', source_path.name)
print('SOURCE_SHEETS', wb.sheetnames)
for ws in wb.worksheets:
    print(f'-- SOURCE_SHEET {ws.title} dimensions={ws.max_row}x{ws.max_column}')
    max_preview_row = min(ws.max_row or 1, 20)
    for row in ws.iter_rows(min_row=1, max_row=max_preview_row, values_only=True):
        vals = list(row)
        while vals and vals[-1] is None:
            vals.pop()
        print(vals)

master = json.loads(master_path.read_text())
values = master.get('values', [])
print(f'-- MASTER_RANGE {master.get("range")} rows={len(values)}')
for i, row in enumerate(values[:5], start=1):
    print('MASTER_ROW', i, row)
headers = values[0] if values else []
for idx, h in enumerate(headers, start=1):
    if h is not None and ('conf' in str(h).lower() or 'name' in str(h).lower() or 'guest' in str(h).lower()):
        print('MASTER_HEADER', idx, repr(h))
