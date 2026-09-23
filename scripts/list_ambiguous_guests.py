import json, re
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path('/home/ubuntu/vegas-sweeps-navigator/work')
source=load_workbook(ROOT/'source_delegate_9_22.xlsx', data_only=True, read_only=False).active
master=json.loads((ROOT/'realnovemberevent_before.json').read_text())['values']
def norm(v):
 s='' if v is None else str(v).upper().strip()
 return re.sub(r'\s+',' ',re.sub(r'[^A-Z0-9]+',' ',s)).strip()
def key(last,first): return f'{norm(first)}|{norm(last)}'
by=defaultdict(list)
for i,row in enumerate(master[1:],2): by[key(row[10],row[9])].append(i)
current=None
for r in range(1,source.max_row+1):
 a=source.cell(r,1).value; b=source.cell(r,2).value; c=source.cell(r,3).value
 if a is not None and c is not None and not str(a).startswith(('THE ','Print ','Group ')) and norm(a)!='LAST NAME':
  current=(r,str(a),str(b or ''),str(c).strip())
 elif current and isinstance(b,str) and b.upper().startswith('ADDL GST:'):
  raw=b.split(':',1)[1].strip(); parts=[p.strip() for p in raw.split(',',1)]
  if len(parts)==2:
   glast,gfirst=parts; candidates=by[key(glast,gfirst)]
   if len(candidates)>1:
    primary_row=by[key(current[1],current[2])]
    print('guest_source_row',r,'guest',gfirst,glast,'primary',current[2],current[1],'conf',current[3],'primary_candidates',primary_row,'guest_candidates',[(x,master[x-1][5],master[x-1][7],master[x-1][19],master[x-1][20]) for x in candidates])
