from __future__ import annotations
import json, re
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path('/home/ubuntu/vegas-sweeps-navigator/work')
source=load_workbook(ROOT/'source_delegate_9_22.xlsx', data_only=True, read_only=False).active
master=json.loads((ROOT/'realnovemberevent_before.json').read_text())['values']

def norm(v):
    if v is None: return ''
    s=str(v).upper().strip()
    s=re.sub(r'[^A-Z0-9]+',' ',s)
    return re.sub(r'\s+',' ',s).strip()

def key(last, first): return f'{norm(first)}|{norm(last)}'

master_by_name=defaultdict(list)
master_guest_by_name=defaultdict(list)
for r_idx,row in enumerate(master[1:], start=2):
    first=row[9] if len(row)>9 else ''
    last=row[10] if len(row)>10 else ''
    if norm(first) or norm(last): master_by_name[key(last,first)].append(r_idx)
    guest=row[61] if len(row)>61 else ''
    if guest: master_guest_by_name[norm(guest)].append(r_idx)

records=[]
current=None
for r in range(1, source.max_row+1):
    a=source.cell(r,1).value; b=source.cell(r,2).value; c=source.cell(r,3).value
    if a is not None and c is not None and not str(a).startswith('THE ') and not str(a).startswith('Print ') and not str(a).startswith('Group ') and norm(a)!='LAST NAME':
        current={'source_row':r,'last':str(a),'first':str(b or ''),'conf':str(c).strip(),'master_rows':master_by_name.get(key(a,b),[]),'guests':[]}
        records.append(current)
    elif current and isinstance(b,str) and b.strip().upper().startswith('ADDL GST:'):
        raw=b.split(':',1)[1].strip()
        parts=[p.strip() for p in raw.split(',',1)]
        if len(parts)==2:
            glast,gfirst=parts
            guest={'source_row':r,'last':glast,'first':gfirst,'master_name_rows':master_by_name.get(key(glast,gfirst),[]),'master_guest_rows':master_guest_by_name.get(norm(f'{gfirst} {glast}'),[])}
            current['guests'].append(guest)

print('source_primary_records',len(records))
print('source_guest_records',sum(len(x['guests']) for x in records))
print('master_name_keys',len(master_by_name))
print('master_guest_name_keys',len(master_guest_by_name))
print('\nUNMATCHED_PRIMARY')
for x in records:
    if len(x['master_rows'])!=1:
        print(x)
print('\nGUEST_MATCH_SUMMARY')
counts=defaultdict(int)
for x in records:
    for g in x['guests']:
        if len(g['master_name_rows'])==1: kind='master_bowler_name'
        elif len(g['master_guest_rows'])==1: kind='master_guest_name'
        elif len(g['master_name_rows'])>1 or len(g['master_guest_rows'])>1: kind='ambiguous'
        else: kind='unmatched'
        counts[kind]+=1
        if kind!='master_bowler_name': print({'primary':x['first']+' '+x['last'],'conf':x['conf'],**g})
print('guest_match_counts',dict(counts))
print('\nDUPLICATE_CONF')
conf_to_rows=defaultdict(list)
for x in records:
    for mr in x['master_rows']:
        conf_to_rows[x['conf']].append(mr)
for conf,rows in conf_to_rows.items():
    if len(rows)>1: print(conf,rows)
