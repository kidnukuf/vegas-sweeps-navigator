from __future__ import annotations
import json,re
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path('/home/ubuntu/vegas-sweeps-navigator/work')
source=load_workbook(ROOT/'source_delegate_9_22.xlsx', data_only=True, read_only=False).active
master=json.loads((ROOT/'realnovemberevent_before.json').read_text())['values']

def norm(v):
 s='' if v is None else str(v).upper().strip()
 return re.sub(r'\s+',' ',re.sub(r'[^A-Z0-9]+',' ',s)).strip()
def key(last,first): return f'{norm(first)}|{norm(last)}'

def parse_source():
 recs=[]; current=None
 for r in range(1,source.max_row+1):
  a=source.cell(r,1).value; b=source.cell(r,2).value; c=source.cell(r,3).value; arr=source.cell(r,5).value; dep=source.cell(r,6).value
  if a is not None and c is not None and not str(a).startswith(('THE ','Print ','Group ')) and norm(a)!='LAST NAME':
   current={'source_row':r,'last':str(a),'first':str(b or ''),'conf':str(c).strip(),'arrival':arr,'departure':dep,'guests':[]}; recs.append(current)
  elif current and isinstance(b,str) and b.strip().upper().startswith('ADDL GST:'):
   raw=b.split(':',1)[1].strip(); parts=[p.strip() for p in raw.split(',',1)]
   if len(parts)==2: current['guests'].append({'source_row':r,'last':parts[0],'first':parts[1]})
 return recs
recs=parse_source(); by=defaultdict(list)
for i,row in enumerate(master[1:],2): by[key(row[10] if len(row)>10 else '',row[9] if len(row)>9 else '')].append(i)
print('primary exact unique',sum(len(by[key(x['last'],x['first'])])==1 for x in recs))
print('primary unmatched',sum(len(by[key(x['last'],x['first'])])==0 for x in recs))
print('primary ambiguous',sum(len(by[key(x['last'],x['first'])])>1 for x in recs))
print('\nAMBIGUOUS_PRIMARY_GROUPS')
seen=set()
for x in recs:
 k=key(x['last'],x['first']); cands=by[k]
 if len(cands)>1 and k not in seen:
  seen.add(k); print('SOURCE',k,[(x2['source_row'],x2['conf'],x2['arrival'],x2['departure']) for x2 in recs if key(x2['last'],x2['first'])==k]); print('MASTER',[(r,master[r-1][5],master[r-1][7],master[r-1][19],master[r-1][20]) for r in cands])
print('\nPRIMARY_CONF_DUPLICATES')
conf=defaultdict(list)
for x in recs: conf[x['conf']].append(x['source_row'])
for c,rs in conf.items():
 if len(rs)>1: print(c,rs)
