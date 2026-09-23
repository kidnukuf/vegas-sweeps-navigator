from __future__ import annotations
import csv, json, re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path('/home/ubuntu/vegas-sweeps-navigator/work')
source=load_workbook(ROOT/'source_delegate_9_22.xlsx', data_only=True, read_only=False).active
master_doc=json.loads((ROOT/'realnovemberevent_before.json').read_text())
master=master_doc['values']

def norm(v):
 s='' if v is None else str(v).upper().strip()
 return re.sub(r'\s+',' ',re.sub(r'[^A-Z0-9]+',' ',s)).strip()
def key(last,first): return f'{norm(first)}|{norm(last)}'
def date_norm(v):
    if isinstance(v,datetime): return v.date().isoformat()
    s='' if v is None else str(v).strip()
    for fmt in ('%m/%d/%Y', '%m/%-d/%Y'):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return s

by_name=defaultdict(list)
for r,row in enumerate(master[1:],2):
 by_name[key(row[10] if len(row)>10 else '', row[9] if len(row)>9 else '')].append(r)

def primary_records():
 recs=[]; cur=None
 for r in range(1,source.max_row+1):
  a=source.cell(r,1).value; b=source.cell(r,2).value; c=source.cell(r,3).value; arr=source.cell(r,5).value; dep=source.cell(r,6).value
  if a is not None and c is not None and not str(a).startswith(('THE ','Print ','Group ')) and norm(a)!='LAST NAME':
   cur={'source_row':r,'last':str(a),'first':str(b or ''),'conf':str(c).strip(),'arrival':date_norm(arr),'departure':date_norm(dep),'guests':[]}; recs.append(cur)
  elif cur and isinstance(b,str) and b.upper().startswith('ADDL GST:'):
   raw=b.split(':',1)[1].strip(); parts=[p.strip() for p in raw.split(',',1)]
   if len(parts)==2: cur['guests'].append({'source_row':r,'last':parts[0],'first':parts[1]})
 return recs
recs=primary_records()
assign=[]; issues=[]
for x in recs:
 cands=by_name[key(x['last'],x['first'])]
 selected=[]; method=''
 if len(cands)==1:
  selected=cands; method='exact-name'
 elif len(cands)>1:
  by_dates=[r for r in cands if date_norm(master[r-1][19] if len(master[r-1])>19 else '')==x['arrival'] and date_norm(master[r-1][20] if len(master[r-1])>20 else '')==x['departure']]
  if len(by_dates)==1:
   selected=by_dates; method='exact-name-and-dates'
  else:
   issues.append({'type':'primary-ambiguous','source_row':x['source_row'],'name':f"{x['first']} {x['last']}",'conf':x['conf'],'candidates':cands,'date_candidates':by_dates})
 else:
  issues.append({'type':'primary-unmatched','source_row':x['source_row'],'name':f"{x['first']} {x['last']}",'conf':x['conf']})
 if selected:
  mr=selected[0]; assign.append({'target_row':mr,'conf':x['conf'],'source_row':x['source_row'],'type':'primary','name':f"{x['first']} {x['last']}",'method':method})
  primary_center=master[mr-1][5] if len(master[mr-1])>5 else ''
 else:
  primary_center=''
 for g in x['guests']:
  gcands=by_name[key(g['last'],g['first'])]
  gsel=[]; gmethod=''
  if len(gcands)==1:
   gsel=gcands; gmethod='exact-name'
  elif len(gcands)>1 and primary_center:
   by_center=[r for r in gcands if norm(master[r-1][5] if len(master[r-1])>5 else '')==norm(primary_center)]
   if len(by_center)==1:
    gsel=by_center; gmethod='exact-name-and-primary-center'
  if len(gsel)==1:
   assign.append({'target_row':gsel[0],'conf':x['conf'],'source_row':g['source_row'],'type':'guest','name':f"{g['first']} {g['last']}",'primary':f"{x['first']} {x['last']}",'method':gmethod})
  else:
   issues.append({'type':'guest-ambiguous' if len(gcands)>1 else 'guest-unmatched','source_row':g['source_row'],'name':f"{g['first']} {g['last']}",'conf':x['conf'],'primary':f"{x['first']} {x['last']}",'candidates':gcands})

by_target=defaultdict(list)
for a in assign: by_target[a['target_row']].append(a)
conflicts=[]
final=[]
for row,items in by_target.items():
 confs=sorted(set(x['conf'] for x in items))
 existing=(master[row-1][18] if len(master[row-1])>18 else '')
 if len(confs)>1:
  conflicts.append({'target_row':row,'existing':existing,'assignments':items})
 elif not existing or str(existing)==confs[0]:
  final.append({'target_row':row,'conf':confs[0],'existing':existing,'assignments':items})
 else:
  conflicts.append({'target_row':row,'existing':existing,'assignments':items})

plan={'source_primary_records':len(recs),'source_guest_records':sum(len(x['guests']) for x in recs),'proposed_target_rows':len(final),'proposed_assignments':len(assign),'issues':issues,'conflicts':conflicts,'final':final}
(ROOT/'confirmation_update_plan.json').write_text(json.dumps(plan,indent=2,default=str))
print('source_primary_records',plan['source_primary_records'])
print('source_guest_records',plan['source_guest_records'])
print('proposed_assignments',plan['proposed_assignments'])
print('proposed_target_rows',plan['proposed_target_rows'])
print('issues',len(issues),'conflicts',len(conflicts))
for kind in sorted(set(x['type'] for x in issues)):
 print(kind,sum(x['type']==kind for x in issues))
print('existing_nonempty_target_rows',sum(bool(x['existing']) for x in final))
