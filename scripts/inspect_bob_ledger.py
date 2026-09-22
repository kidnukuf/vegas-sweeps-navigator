from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/upload/BOB2026ledgerMASTER.xlsx')
wb = load_workbook(path, read_only=True, data_only=False)
print(f'workbook: {path.name}')
print(f'sheets: {wb.sheetnames}')
for ws in wb.worksheets:
    print(f'--- sheet: {ws.title!r} ---')
    print(f'dimensions: {ws.max_row} rows x {ws.max_column} columns')
    nonempty_rows = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if any(v not in (None, '') for v in vals):
            nonempty_rows.append(vals)
    print(f'nonempty_rows: {len(nonempty_rows)}')
    for i, vals in enumerate(nonempty_rows[:8], 1):
        print(f'{i}: {vals}')
    print()
