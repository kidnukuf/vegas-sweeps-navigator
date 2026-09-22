from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import csv
from openpyxl import load_workbook

SOURCE = Path('/home/ubuntu/upload/BOB2026ledgerMASTER.xlsx')
OUT_TSV = Path('/home/ubuntu/upload/realnovemberevent_import.tsv')
OUT_CSV = Path('/home/ubuntu/upload/realnovemberevent_import.csv')

TARGET_HEADERS = [
    'Bowler ID','Phone','Email','Squad Day & Time','Lane #','Center','Coordinator','Team #','Captain','First Name','Last Name','Under 21?','Sanction #','# Games','Best Avg','Team Name','League Member','T-Shirt Size','Hotel Confirmation','Check In','Check Out','Roommate First Name','Roommate Last Name','2nd Squad Time','Lane #','Pool QR','Pool Used','Banquet QR','Banquet Used','#A Pool QR','#A Pool Used','#A Banquet QR','#A Banquet Used','#B Pool QR','#B Pool Used','#B Banquet QR','#B Banquet Used','2nd Banquet QR','2nd Banquet Used','2nd Pool QR','2nd Pool Used','Q1 Overall Experience?','Q1 answer','Q2 Bowling Venue?','Q2 Answer','Q3 Event Organization?','Q3 Answer','Q4 Pool Party? (If applicable)','Q4 Answer','Q5 Banquet Experience?','Q5 Answer','Q6 This App?','Q6 Answer','Q7 League App Interest?','Q7 Answer','Q8 Additional Comments or Concerns','Q8 Answer','Q9 Testimonial Permission?','Q9 Answer','Q10 Attend Next Year?','Q10 Answer','Guest Name','Additional Guest','Claim Code','Bill Breakdown','Team Score','Hotel Room Id','Seating Arrangement'
]


def text(v: object) -> str:
    if v is None:
        return ''
    if isinstance(v, datetime):
        return f'{v.month}/{v.day}/{v.year}'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, str) and v.startswith('='):
        return ''
    return str(v).strip()


def key(center: object, team: object, first: object, last: object) -> tuple[str, str, str, str]:
    return (text(center).lower(), text(team).lower(), text(first).lower(), text(last).lower())


def read_rows(ws, start_row: int, fields: dict[str, int]):
    out = []
    for row_no, row in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start_row):
        def get(name):
            i = fields.get(name)
            return row[i] if i is not None and i < len(row) else None
        center, team, first, last = get('center'), get('team'), get('first'), get('last')
        if not (text(center) and text(first) and text(last)):
            continue
        out.append({'row': row_no, **{name: get(name) for name in fields}})
    return out


wb = load_workbook(SOURCE, read_only=True, data_only=True)
master_ws = wb['Master']
room_ws = wb['Room list']
lanes_ws = wb['Lanes by capt']

# Master header rows are split across rows 3-4; all indices are zero-based.
master_fields = {
    'center': 0, 'squad': 1, 'lane': 2, 'team': 3, 'captain': 4,
    'first': 5, 'last': 6, 'under21': 8, 'sanction': 9, 'games': 10,
    'bestavg': 17, 'teamname': 18, 'leagueMember': 19, 'hotelConf': 24,
    'checkin': 25, 'checkout': 26, 'roomType': 27, 'roomWith': 29,
    'roomFirst': 30, 'roomLast': 31, 'roomAmount': 32, 'cocktail': 33,
    'totalDue': 34, 'notes': 35, 'phone': 36, 'email': 37,
}
room_fields = {
    'center': 0, 'team': 1, 'first': 2, 'last': 3, 'checkin': 5,
    'checkout': 6, 'roomType': 7, 'special': 8, 'roomWith': 9,
    'roomFirst': 10, 'roomLast': 11, 'roomAmount': 12, 'cocktail': 13,
    'notes': 15, 'phone': 16, 'email': 17, 'teamname': 18,
}
lane_fields = {'lane': 0, 'center': 1, 'team': 2, 'teamname': 3, 'first': 4, 'last': 5}

master_rows = read_rows(master_ws, 5, master_fields)
room_rows = read_rows(room_ws, 5, room_fields)
lane_rows = read_rows(lanes_ws, 4, lane_fields)
room_map = {key(r['center'], r['team'], r['first'], r['last']): r for r in room_rows}
lane_map = {key(r['center'], r['team'], r['first'], r['last']): r for r in lane_rows}

# Keep the Master sheet as the authoritative roster. Supplement only blanks from the other tabs.
rows = []
for source in master_rows:
    k = key(source['center'], source['team'], source['first'], source['last'])
    room = room_map.get(k, {})
    lane = lane_map.get(k, {})
    def pick(field):
        value = source.get(field)
        return value if text(value) else room.get(field, lane.get(field, ''))
    row = [''] * len(TARGET_HEADERS)
    # App-managed fields intentionally remain blank.
    values = {
        'Phone': pick('phone'), 'Email': pick('email'), 'Squad Day & Time': pick('squad'),
        'Lane #': pick('lane'), 'Center': pick('center'), 'Team #': pick('team'),
        'Captain': pick('captain'), 'First Name': pick('first'), 'Last Name': pick('last'),
        'Under 21?': pick('under21'), 'Sanction #': pick('sanction'), '# Games': pick('games'),
        'Best Avg': pick('bestavg'), 'Team Name': pick('teamname'), 'League Member': pick('leagueMember'),
        'Hotel Confirmation': pick('hotelConf'), 'Check In': pick('checkin'), 'Check Out': pick('checkout'),
        'Roommate First Name': pick('roomFirst'), 'Roommate Last Name': pick('roomLast'),
    }
    for header, value in values.items():
        row[TARGET_HEADERS.index(header)] = text(value)
    rows.append(row)

# Remove accidental duplicate source entries by exact center/team/name identity, preserving first source row.
deduped = []
seen = set()
for row in rows:
    k = key(row[5], row[7], row[9], row[10])
    if k in seen:
        continue
    seen.add(k)
    deduped.append(row)
rows = deduped

with OUT_TSV.open('w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f, delimiter='\t', lineterminator='\n')
    writer.writerows(rows)
with OUT_CSV.open('w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f, lineterminator='\n')
    writer.writerow(TARGET_HEADERS)
    writer.writerows(rows)

print('source_master_rows', len(master_rows))
print('source_room_rows', len(room_rows))
print('source_lane_rows', len(lane_rows))
print('output_rows', len(rows))
print('output_columns', len(TARGET_HEADERS))
print('supplemented_room_matches', sum(1 for r in master_rows if key(r['center'],r['team'],r['first'],r['last']) in room_map))
print('supplemented_lane_matches', sum(1 for r in master_rows if key(r['center'],r['team'],r['first'],r['last']) in lane_map))
print('first_data_row', rows[0][:24] if rows else [])
print('last_data_row', rows[-1][:24] if rows else [])
print('tsv', OUT_TSV)
print('csv', OUT_CSV)
