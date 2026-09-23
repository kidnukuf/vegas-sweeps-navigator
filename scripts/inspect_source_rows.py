from pathlib import Path
from openpyxl import load_workbook

p=Path('/home/ubuntu/vegas-sweeps-navigator/work/source_delegate_9_22.xlsx')
wb=load_workbook(p, data_only=True, read_only=False)
ws=wb[wb.sheetnames[0]]
print('sheet', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
for r in range(1, min(ws.max_row, 80)+1):
    vals=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
    populated=[f'{c}:{v!r}' for c,v in enumerate(vals,1) if v is not None]
    print(f'ROW {r}: ' + ' | '.join(populated))
